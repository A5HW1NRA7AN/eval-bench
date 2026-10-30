"""`bench export`: every stage's output as one Excel workbook (data/bench.xlsx), a sheet per table.

Sheets that need a stage not run yet are left out. Each sheet has a frozen header, filters and wrapped text.
"""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from bench.checks import RESULTS, all_checks, reason_group
from bench.config import DATA_DIR
from bench.evaluate import KS, summary
from bench.facts import FACTS
from bench.llm import TRACE
from bench.pdf import PAGES, read_jsonl
from bench.questions import QUESTIONS
from bench.segments import SEGMENTS
from bench.testbed import CHUNKS

WORKBOOK = DATA_DIR / "bench.xlsx"
CELL_LIMIT = 32_000  # Excel holds at most 32,767 characters in a cell
WIDE = 60  # width of text columns; others fit their content
HEADER = PatternFill("solid", fgColor="DDE4EE")
BAD = PatternFill("solid", fgColor="F8D7DA")


def _rows(path: Path) -> list[dict[str, Any]]:
    return read_jsonl(path) if path.exists() else []


def _cell(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    return ILLEGAL_CHARACTERS_RE.sub("", value[:CELL_LIMIT]) if isinstance(value, str) else value  # PDFs hold control characters


def _sheet(book: Workbook, title: str, rows: list[dict[str, Any]], bad: Callable[[dict], bool] | None = None) -> None:
    """One sheet from a list of dicts; rows for which `bad` is true are shaded red."""
    if not rows:
        return
    ws = book.create_sheet(title)
    head = list(rows[0])
    ws.append(head)
    for row in rows:
        ws.append([_cell(row.get(k)) for k in head])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
            if bad and bad(row):
                cell.fill = BAD
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), HEADER
    for col, name in zip(ws.columns, head):
        longest = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(len(name), longest) + 2, WIDE)
        if longest > WIDE:
            for c in col[1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes, ws.auto_filter.ref = "B2", ws.dimensions


def run_export() -> Path:
    segments, facts, questions = _rows(SEGMENTS), _rows(FACTS), _rows(QUESTIONS)
    results, trace = _rows(RESULTS), _rows(TRACE)
    chunks = {c["chunk_id"]: c for c in _rows(CHUNKS)}
    facts_by, questions_by = {r["segment_id"]: r for r in facts}, {r["segment_id"]: r for r in questions}
    scored = {r["question_id"]: r for r in results}
    answerable = [r for r in results if r.get("answerable", True)]

    def missed(r: dict[str, Any]) -> bool:  # an answerable question whose evidence wasn't in the top 10: shaded red
        return bool(r.get("answerable", True) and r.get("hit@10") == 0)

    def passage(cid: str) -> str:
        c = chunks.get(cid)
        return f"{c['file_name']} p{c['pages'][0]}: {' '.join(c['text'].split())[:300]}" if c else cid

    book = Workbook()
    book.remove(book.active)
    _sheet(book, "Overview", [
        {"stage": "pages", "rows": len(_rows(PAGES)), "detail": "cleaned PDF pages"},
        {"stage": "segments", "rows": len(segments), "detail": f"{sum(s['use_for_questions'] for s in segments)} used for questions"},
        {"stage": "facts", "rows": sum(len(r["facts"]) for r in facts),
         "detail": f"{sum(len(r['rejected']) for r in facts)} rejected by the audit"},
        {"stage": "questions", "rows": sum(len(r["questions"]) for r in questions),
         "detail": f"{sum(len(r['rejected']) for r in questions)} rejected by the audit"},
        {"stage": "test bed", "rows": len(chunks), "detail": "chunks in Qdrant"},
        {"stage": "evaluate", "rows": len(results),
         "detail": ", ".join(f"{m} {v:.2f}" for m, v in summary(answerable, "type")[0].items() if m.startswith(("hit", "rr")))
         + f"; {len(results) - len(answerable)} unanswerable" if answerable else "not run"},
        {"stage": "agent calls", "rows": len(trace), "detail": f"{sum(t['seconds'] for t in trace) / 60:.1f} min in data/trace.jsonl"},
    ])
    _sheet(book, "Segments", [
        {"segment_id": s["segment_id"], "title": s["title"], "pdf": s["file_name"], "pages": f"{s['pages'][0]}-{s['pages'][-1]}" if s["pages"] else "",
         "used_for_questions": s["use_for_questions"], "chars": len(s["text"]),
         "facts": len(facts_by.get(s["segment_id"], {}).get("facts", [])),
         "facts_rejected": len(facts_by.get(s["segment_id"], {}).get("rejected", [])),
         "questions": len(questions_by.get(s["segment_id"], {}).get("questions", [])),
         "questions_rejected": len(questions_by.get(s["segment_id"], {}).get("rejected", []))} for s in segments])
    _sheet(book, "Facts", [{"fact_id": f["fact_id"], "size": f.get("size"), "claim": f["claim"], "quote": f["quote"],
                            "pages": f"{f['pages'][0]}-{f['pages'][-1]}", "chars": len(f["quote"]), "topic": f.get("topic"),
                            "review": f.get("review", {}).get("verdict"), "claim_before_review": f.get("review", {}).get("original"),
                            "review_reason": f.get("review", {}).get("reason"), "model": r.get("model"), "version": r["version"]}
                           for r in facts for f in r["facts"]])
    _sheet(book, "Rejected facts", [{"segment_id": r["segment_id"], "reason": f["reason"], "claim": f["claim"], "quote": f.get("quote"),
                                     "sentences": f"{f['first']}–{f['last']}" if "first" in f else None, "topic": f.get("topic")}
                                    for r in facts for f in r["rejected"]])
    _sheet(book, "Questions", [
        {"question_id": q["question_id"], "type": q["type"], "question": q["question"], "answer": q.get("answer"),
         "absent_term": q.get("absent_term") or None, "rewrites": q.get("source_question"),
         "review": q.get("review", {}).get("verdict"), "question_before_review": q.get("review", {}).get("original"),
         "fact_ids": ", ".join(q["fact_ids"]), "evidence": "\n---\n".join(ev["quote"] for ev in q["evidence"]),
         "first_rank": scored.get(q["question_id"], {}).get("first_rank"),
         **{f"hit@{k}": scored.get(q["question_id"], {}).get(f"hit@{k}") for k in KS}}
        for r in questions for q in r["questions"]], bad=lambda row: missed(scored.get(row["question_id"], {})))
    _sheet(book, "Rejected questions", [{"segment_id": r["segment_id"], "reason": q["reason"], "type": q["type"],
                                         "question": q["question"], "answer": q.get("answer"), "fact_ids": ", ".join(q["fact_ids"])}
                                        for r in questions for q in r["rejected"]])
    _sheet(book, "Results", [
        {"question_id": r["question_id"], "type": r["type"], "question": r["question"], "answerable": r.get("answerable", True),
         "first_rank": r.get("first_rank"), **{f"hit@{k}": r.get(f"hit@{k}") for k in KS}, "rr": r.get("rr"),
         "all_facts@10": r.get("all_facts@10"), "top_score": r["top"][0][1] if r["top"] else None,
         "expected": "\n---\n".join(r["expected"]),
         **{f"top{n}": f"{score:.3f} {passage(cid)}" for n, (cid, score) in enumerate(r["top"][:3], start=1)}}
        for r in results], bad=missed)
    if answerable:
        pdf_of = {s["segment_id"]: s["file_name"] for s in segments}
        _sheet(book, "Metrics by type", summary(answerable, "type"))
        _sheet(book, "Metrics by PDF", summary([{**r, "pdf": pdf_of.get(r["question_id"].rsplit(":", 1)[0], "several PDFs")}
                                                for r in answerable], "pdf"))
    rejects = Counter((stage, reason_group(item["reason"])) for stage, rows in (("facts", facts), ("questions", questions))
                      for r in rows for item in r["rejected"])
    _sheet(book, "Audit reasons", [{"stage": s, "reason": reason, "count": n} for (s, reason), n in rejects.most_common()])
    _sheet(book, "Checks", [{"stage": stage, "check": name, "ok": ok, "detail": detail}
                            for stage, checks in all_checks().items() for name, ok, detail in checks], bad=lambda row: not row["ok"])
    _sheet(book, "Agent calls", [{"time": t["time"], "role": t["role"], "label": t["label"], "model": t["model"],
                                  "seconds": t["seconds"], "input_tokens": t["input_tokens"], "output_tokens": t["output_tokens"],
                                  "output": t["output"], "prompt": t["prompt"]} for t in trace])
    book.save(WORKBOOK)
    print(f"wrote {WORKBOOK}: {', '.join(book.sheetnames)}")
    return WORKBOOK
