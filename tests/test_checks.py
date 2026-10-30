import json

import pytest

from bench import checks, serve

TEXT = "Give the calf colostrum within two hours of birth."


def _write(path, rows):
    path.write_text("\n".join(json.dumps(r) for r in rows))


def _data(tmp_path, monkeypatch, fact_start=0):
    (tmp_path / "pdfs").mkdir()
    (tmp_path / "pdfs" / "book.pdf").write_bytes(b"")
    pages = [{"file_name": "book.pdf", "page": n, "text": TEXT, "raw_text": TEXT, "garbled": False, "printed": n} for n in (1, 2, 3)]
    segment = {"segment_id": "book-01", "file_name": "book.pdf", "title": "1 Calves", "pages": [1, 2], "use_for_questions": True,
               "headings": [], "text": f"{TEXT}\n\n{TEXT}",
               "page_spans": [{"page": 1, "start": 0, "end": len(TEXT)}, {"page": 2, "start": len(TEXT) + 2, "end": 2 * len(TEXT) + 2}]}
    fact = {"fact_id": "book-01:f01", "claim": "Calves need colostrum early.", "quote": "colostrum within two hours",
            "start": fact_start, "end": fact_start + 26, "pages": [1, 1]}
    paths = {"PAGES": tmp_path / "pages.jsonl", "SEGMENTS": tmp_path / "segments.jsonl", "FACTS": tmp_path / "facts.jsonl"}
    _write(paths["PAGES"], pages)
    _write(paths["SEGMENTS"], [segment])
    _write(paths["FACTS"], [{"segment_id": "book-01", "version": checks.FACTS_VERSION, "facts": [fact], "rejected": []}])
    for name, path in paths.items():
        monkeypatch.setattr(checks, name, path)
    monkeypatch.setattr(checks, "PDF_DIR", tmp_path / "pdfs")
    monkeypatch.setattr(checks, "QUESTIONS", tmp_path / "questions.jsonl")  # not written: never the real data/ file
    monkeypatch.setattr(checks, "DATA_DIR", tmp_path)
    return paths


def test_checks_catch_a_lost_page_and_a_moved_quote(tmp_path, monkeypatch):
    _data(tmp_path, monkeypatch, fact_start=3)  # the quote really starts at TEXT.index("colostrum")
    results = {name: ok for name, ok, _ in checks.check_segment() + checks.check_facts()}
    assert results["every PDF in pdfs/ is extracted"] and results["page positions point at each page's text"]
    assert not results["every page with text is in exactly one segment"]  # page 3 is in no segment
    assert not results["every quote is verbatim at its recorded position"]
    fixed = tmp_path / "fixed"
    fixed.mkdir()
    _data(fixed, monkeypatch, fact_start=TEXT.index("colostrum"))
    assert dict((n, ok) for n, ok, _ in checks.check_facts())["every quote is verbatim at its recorded position"]


def test_api_routes(tmp_path, monkeypatch):
    paths = _data(tmp_path, monkeypatch)
    monkeypatch.setattr(serve, "SEGMENTS", paths["SEGMENTS"])
    monkeypatch.setattr(serve, "FACTS", tmp_path / "missing.jsonl")
    assert serve.route("/", {})[0] == 200
    status, segments = serve.route("/segments", {"doc": "book"})
    assert status == 200 and segments[0]["segment_id"] == "book-01" and "text" not in segments[0]
    assert serve.route("/segments/book-01", {})[1]["text"].startswith("Give the calf")
    assert serve.route("/segments/book-99", {})[0] == 404 and serve.route("/nope", {})[0] == 404
    with pytest.raises(LookupError):  # the server answers 404: run `bench facts` first
        serve.route("/facts", {})


def test_agent_calls_are_traced_and_the_workbook_has_every_sheet(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from pydantic import BaseModel
    from pydantic_ai import Agent, ModelResponse, NativeOutput, TextPart
    from pydantic_ai.models.function import FunctionModel

    from bench import export, llm

    class Out(BaseModel):
        items: list[str]

    agent = Agent(FunctionModel(lambda m, i: ModelResponse(parts=[TextPart('{"items": ["a", "b"]}')])),
                  output_type=NativeOutput(Out, strict=True))
    assert llm.call("tester", agent, "say two things", "book-01").items == ["a", "b"]
    call = json.loads(llm.TRACE.read_text())
    assert call["role"] == "tester" and call["prompt"] == "say two things" and call["output"] == {"items": ["a", "b"]}

    paths = _data(tmp_path, monkeypatch, fact_start=TEXT.index("colostrum"))
    for name in ("SEGMENTS", "FACTS", "PAGES"):
        monkeypatch.setattr(export, name, paths[name])
    for name, path in {"QUESTIONS": "q.jsonl", "RESULTS": "r.jsonl", "CHUNKS": "c.jsonl", "WORKBOOK": "bench.xlsx"}.items():
        monkeypatch.setattr(export, name, tmp_path / path)
    monkeypatch.setattr(export, "TRACE", llm.TRACE)
    book = load_workbook(export.run_export())
    assert book.sheetnames == ["Overview", "Segments", "Facts", "Checks", "Agent calls"]  # no questions or results yet
    row = {head.value: cell.value for head, cell in zip(book["Facts"][1], book["Facts"][2])}
    assert row["fact_id"] == "book-01:f01" and row["claim"] == "Calves need colostrum early."
